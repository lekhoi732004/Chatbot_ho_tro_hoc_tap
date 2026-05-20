"""
DOCX reader tool.
Extracts text (paragraphs, headings) and optionally tables from .docx files.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from src.utils.config_loader import get_config
from src.utils.logger import get_logger

logger = get_logger("docx_reader")


def extract_docx(
    file_path: str,
    extract_tables: Optional[bool] = None,
) -> Tuple[str, Dict]:
    """
    Extract content from a .docx file.

    Args:
        file_path: Path to the .docx file.
        extract_tables: Override config setting for table extraction.

    Returns:
        (full_text, structured_data)
        - full_text: All text joined by newlines.
        - structured_data: {paragraphs, headings, tables}.
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx not installed. Run: pip install python-docx")

    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    cfg = get_config()
    do_extract_tables = (
        extract_tables
        if extract_tables is not None
        else cfg.get("tools.docx.extract_tables", True)
    )

    logger.info(f"Extracting DOCX: {file_path}")

    doc = Document(file_path)

    paragraphs = []
    headings = []
    full_lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading ", "").strip()
            headings.append({"level": level, "text": text})
            full_lines.append(f"\n## {text}\n")
        else:
            paragraphs.append(text)
            full_lines.append(text)

    tables_data = []
    if do_extract_tables:
        for tbl in doc.tables:
            rows = []
            for row in tbl.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                tables_data.append(rows)
                # Append simple text representation
                full_lines.append("\n[TABLE]\n" + _table_to_text(rows))

    full_text = "\n".join(full_lines).strip()

    structured_data = {
        "paragraphs": paragraphs,
        "headings": headings,
        "tables": tables_data,
        "file_name": Path(file_path).name,
    }

    logger.info(
        f"Extracted {len(paragraphs)} paragraphs, {len(headings)} headings, "
        f"{len(tables_data)} tables from {Path(file_path).name}"
    )
    return full_text, structured_data


def _table_to_text(rows: List[List[str]]) -> str:
    """Convert table rows to a markdown-like text representation."""
    if not rows:
        return ""
    lines = []
    header = rows[0]
    lines.append(" | ".join(header))
    lines.append("-" * 40)
    for row in rows[1:]:
        lines.append(" | ".join(row))
    return "\n".join(lines)


def get_docx_metadata(file_path: str) -> Dict:
    """Return basic metadata from a DOCX file."""
    try:
        from docx import Document
        doc = Document(file_path)
        props = doc.core_properties
        return {
            "title": props.title or Path(file_path).stem,
            "author": props.author or "Unknown",
            "created": str(props.created),
            "modified": str(props.modified),
            "file_size_kb": round(os.path.getsize(file_path) / 1024, 2),
        }
    except Exception:
        return {
            "title": Path(file_path).stem,
            "author": "Unknown",
            "file_size_kb": round(os.path.getsize(file_path) / 1024, 2),
        }
